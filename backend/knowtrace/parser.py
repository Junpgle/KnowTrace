"""Parse text-based PDF/PPTX into page- or slide-attributed chunks.

This module deliberately does not OCR scans or interpret diagrams; missing
text is reported rather than hallucinated.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz
from pptx import Presentation

SUPPORTED_EXTENSIONS = {".pdf", ".pptx"}


class ParseError(ValueError):
    pass


def _pages(filename: str, data: bytes) -> list[str]:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError("Only PDF and PPTX files are currently supported")
    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ParseError("File content is not a PDF")
        try:
            with fitz.open(stream=data, filetype="pdf") as document:
                return [page.get_text(sort=True).strip() for page in document]
        except Exception as exc:
            raise ParseError("Unable to read PDF") from exc
    if not data.startswith(b"PK"):
        raise ParseError("File content is not a PPTX")
    try:
        presentation = Presentation(BytesIO(data))
        result = []
        for slide in presentation.slides:
            items: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        items.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        items.append(" | ".join(cell.text.strip() for cell in row.cells))
            result.append("\n".join(items))
        return result
    except Exception as exc:
        raise ParseError("Unable to read PPTX") from exc


def extract_chunks(filename: str, data: bytes, *, chunk_size: int = 1600,
                   overlap: int = 150) -> tuple[int, list[dict]]:
    """Return physical page count and chunks with 1-indexed page references."""
    if chunk_size <= overlap or overlap < 0:
        raise ValueError("chunk_size must be greater than overlap")
    pages = _pages(filename, data)
    chunks: list[dict] = []
    for number, text in enumerate(pages, start=1):
        if not text:
            continue
        for start in range(0, len(text), chunk_size - overlap):
            section = text[start:start + chunk_size].strip()
            if section:
                chunks.append({"page_start": number, "page_end": number,
                               "content_kind": "text", "content_text": section})
            if start + chunk_size >= len(text):
                break
    if not pages:
        raise ParseError("Document has no pages")
    if not chunks:
        if Path(filename).suffix.lower() == ".pdf":
            # A scanned or image-only PDF is valid source material. It can be
            # rendered page-by-page for explicit Ollama/MiMo multimodal analysis.
            # A zero-text document must not be misrepresented as parsed text.
            return len(pages), []
        raise ParseError("PPTX contains no extractable text; image-only slides are not yet supported")
    return len(pages), chunks
