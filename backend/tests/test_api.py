"""API regression tests on an isolated temporary SQLite database."""
from io import BytesIO

import fitz
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from knowtrace.bkt import BKTParameters, update_mastery
from knowtrace.main import create_app


@pytest.fixture
def client(tmp_path):
    app = create_app(db_path=tmp_path / "test.sqlite3", storage_dir=tmp_path / "uploads")
    with TestClient(app) as instance:
        yield instance


def make_course(client, title="Computer Networks"):
    response = client.post("/api/courses", json={"title": title})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def test_health_and_courses(client):
    assert client.get("/health").json()["status"] == "ok"
    cid = make_course(client)
    assert client.get("/api/courses").json()[0]["id"] == cid
    assert client.get(f"/api/courses/{cid}").json()["title"] == "Computer Networks"
    assert client.get("/api/courses/unknown").status_code == 404


def test_chapter_and_knowledge_point_course_integrity(client):
    cid, other = make_course(client), make_course(client, "Other")
    chapter = client.post(f"/api/courses/{cid}/chapters", json={"title": "Transport", "ordinal": 7}).json()
    assert chapter["id"]
    bad = client.post(f"/api/courses/{other}/knowledge-points", json={"canonical_name": "TCP", "chapter_id": chapter["id"]})
    assert bad.status_code == 422
    good = client.post(f"/api/courses/{cid}/knowledge-points", json={"canonical_name": "TCP", "chapter_id": chapter["id"]})
    assert good.status_code == 201
    assert client.get(f"/api/courses/{cid}/knowledge-points").json()[0]["canonical_name"] == "TCP"


def test_pdf_upload_and_page_provenance(client):
    cid = make_course(client)
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((70, 70), "TCP handshake relies on sequence numbers")
    page = pdf.new_page()
    page.insert_text((70, 70), "UDP is a transport protocol")
    content = pdf.tobytes()
    response = client.post(f"/api/courses/{cid}/documents", files={"file": ("lecture.pdf", content, "application/pdf")})
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["total_pages"] == 2 and data["chunks"] == 2
    chunks = client.get(f"/api/documents/{data['id']}/chunks").json()
    assert chunks[0]["page_start"] == 1 and "TCP handshake" in chunks[0]["content_text"]
    assert chunks[1]["page_start"] == 2 and "UDP" in chunks[1]["content_text"]


def test_pptx_upload(client):
    cid = make_course(client)
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[5])
    slide.shapes.title.text = "OSI seven-layer architecture"
    out = BytesIO()
    pres.save(out)
    response = client.post(f"/api/courses/{cid}/documents", files={"file": ("network.pptx", out.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert response.status_code == 201, response.text
    assert response.json()["chunks"] >= 1


def test_reject_unsupported_and_fake_pdf(client):
    cid = make_course(client)
    bad_ext = client.post(f"/api/courses/{cid}/documents", files={"file": ("sample.txt", b"text", "text/plain")})
    assert bad_ext.status_code == 415
    invalid = client.post(f"/api/courses/{cid}/documents", files={"file": ("sample.pdf", b"fake", "application/pdf")})
    assert invalid.status_code == 422


def test_bkt_update():
    params = BKTParameters()
    posterior_good, after_good = update_mastery(0.2, True, params)
    posterior_bad, after_bad = update_mastery(0.2, False, params)
    assert 0 < posterior_bad < posterior_good < 1
    assert after_good > posterior_good and after_bad > posterior_bad
    with pytest.raises(ValueError):
        update_mastery(1.1, True)
